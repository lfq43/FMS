from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "GitLab协作开发从零开始指南.md"
OUTPUT = ROOT / "GitLab协作开发从零开始指南.docx"

FONT_EAST_ASIA = "微软雅黑"
FONT_LATIN = "Calibri"
FONT_MONO = "Consolas"


def set_run_font(run, name=FONT_LATIN, east_asia=FONT_EAST_ASIA, size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_style_font(style, name=FONT_LATIN, east_asia=FONT_EAST_ASIA, size=None, color=None, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def set_paragraph_spacing(style, before=0, after=6, line=1.25):
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def set_cell_shading(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)

    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color="666666")


def add_hyperlink(paragraph, url, text=None):
    text = text or url
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_inline_runs(paragraph, text):
    url_pattern = re.compile(r"(https?://[^\s>]+)")
    bold_pattern = re.compile(r"\*\*(.+?)\*\*")
    pos = 0
    for match in url_pattern.finditer(text):
        before = text[pos : match.start()]
        add_bold_runs(paragraph, before, bold_pattern)
        add_hyperlink(paragraph, match.group(1))
        pos = match.end()
    add_bold_runs(paragraph, text[pos:], bold_pattern)


def add_bold_runs(paragraph, text, bold_pattern):
    pos = 0
    for match in bold_pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=11)
        run = paragraph.add_run(match.group(1))
        set_run_font(run, size=11, bold=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=11)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=11)
    set_paragraph_spacing(normal, before=0, after=6, line=1.25)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[style_name]
        set_style_font(style, size=size, color=color, bold=True)
        set_paragraph_spacing(style, before=before, after=after, line=1.25)

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        set_style_font(style, size=11)
        set_paragraph_spacing(style, before=0, after=4, line=1.25)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)

    code_style = doc.styles.add_style("CodeBlock", 1)
    set_style_font(code_style, name=FONT_MONO, east_asia="等线", size=9.5, color="1F2937")
    set_paragraph_spacing(code_style, before=3, after=3, line=1.1)
    code_style.paragraph_format.left_indent = Inches(0.15)
    code_style.paragraph_format.right_indent = Inches(0.15)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def build():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    style_document(doc)

    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    set_run_font(run, size=22, color="0B2545", bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("适用项目：")
    set_run_font(run, size=10.5, color="555555", bold=True)
    add_hyperlink(subtitle, "https://jihulab.com/lfq43-group/FMS")

    in_code = False
    first_line = True
    for raw in lines[1:]:
        line = raw.rstrip()
        if first_line and (not line or line.startswith("适用项目")):
            if line.startswith("适用项目"):
                first_line = False
            continue
        first_line = False

        if line.startswith("```"):
            in_code = not in_code
            continue

        if in_code:
            p = doc.add_paragraph(style="CodeBlock")
            set_cell_shading(p, "F3F4F6")
            run = p.add_run(line if line else " ")
            set_run_font(run, name=FONT_MONO, east_asia="等线", size=9.5, color="111827")
            continue

        if not line:
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
            p = doc.add_paragraph(style=style)
            add_inline_runs(p, text)
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        bullet = re.match(r"^\s*-\s+(.+)$", line)

        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, numbered.group(2))
        elif bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, bullet.group(1))
        else:
            p = doc.add_paragraph()
            add_inline_runs(p, line)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
